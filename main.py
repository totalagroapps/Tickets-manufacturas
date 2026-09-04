from fastapi import FastAPI, Depends, Request, Form, status, HTTPException, UploadFile, File, Response
from fastapi.responses import HTMLResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from sqlalchemy.orm import Session
from sqlalchemy import text
from datetime import datetime, timedelta
from jose import jwt, JWTError
from passlib.context import CryptContext
from typing import Optional, List
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

import models
from database import engine, get_db, SessionLocal

from contextlib import asynccontextmanager

def create_default_users():
    db = SessionLocal()
    try:
        # Check if client user exists
        client_user = db.query(models.User).filter(models.User.username == "cliente").first()
        if not client_user:
            client_user = models.User(
                username="cliente",
                password_hash=get_password_hash("1234"),
                role=models.UserRole.CLIENT
            )
            db.add(client_user)

        # Check if admin user exists
        admin_user = db.query(models.User).filter(models.User.username == "admin").first()
        if not admin_user:
            admin_user = models.User(
                username="admin",
                password_hash=get_password_hash("admin123"),
                role=models.UserRole.ADMIN
            )
            db.add(admin_user)

        db.commit()
    finally:
        db.close()

@asynccontextmanager
async def lifespan(app: FastAPI):
    # Startup
    models.Base.metadata.create_all(bind=engine)
    
    # Intento de agregar columnas nuevas si ya existía la tabla (útil para SQLite y Postgres sin alembic)
    try:
        with engine.connect() as conn:
            conn.execute(text("ALTER TABLE tickets ADD COLUMN responsable VARCHAR DEFAULT 'Soporte'"))
            conn.commit()
    except Exception:
        pass # La columna ya existe

    try:
        with engine.connect() as conn:
            conn.execute(text("ALTER TABLE tickets ADD COLUMN celular_solicitante VARCHAR"))
            conn.commit()
    except Exception:
        pass # La columna ya existe
        
    create_default_users()
    yield
    # Shutdown
    pass

app = FastAPI(lifespan=lifespan)

# Para archivos estáticos si los hay (opcional)
import os
os.makedirs("static", exist_ok=True)
os.makedirs("templates", exist_ok=True)
app.mount("/static", StaticFiles(directory="static"), name="static")

templates = Jinja2Templates(directory="templates")

pwd_context = CryptContext(schemes=["bcrypt"], deprecated="auto")

SECRET_KEY = "super_secret_key_change_in_production"
ALGORITHM = "HS256"

def verify_password(plain_password, hashed_password):
    return pwd_context.verify(plain_password, hashed_password)

def get_password_hash(password):
    return pwd_context.hash(password)

def create_access_token(data: dict):
    to_encode = data.copy()
    expire = datetime.utcnow() + timedelta(hours=24)
    to_encode.update({"exp": expire})
    encoded_jwt = jwt.encode(to_encode, SECRET_KEY, algorithm=ALGORITHM)
    return encoded_jwt

def get_current_user(request: Request, db: Session = Depends(get_db)):
    token = request.cookies.get("access_token")
    if not token:
        return None
    try:
        # Check for Bearer prefix and strip it if exists, but we are setting just the token
        if token.startswith("Bearer "):
            token = token[7:]
        payload = jwt.decode(token, SECRET_KEY, algorithms=[ALGORITHM])
        username: str = payload.get("sub")
        if username is None:
            return None
        user = db.query(models.User).filter(models.User.username == username).first()
        return user
    except JWTError:
        return None

def get_current_admin(user: models.User = Depends(get_current_user)):
    if not user or user.role != models.UserRole.ADMIN:
        return None
    return user

@app.get("/", response_class=HTMLResponse)
def root_redirect(request: Request, user: models.User = Depends(get_current_user)):
    if not user:
        return RedirectResponse(url="/login", status_code=status.HTTP_302_FOUND)
    if user.role == models.UserRole.ADMIN:
        return RedirectResponse(url="/admin", status_code=status.HTTP_302_FOUND)
    return RedirectResponse(url="/tickets", status_code=status.HTTP_302_FOUND)

@app.get("/login", response_class=HTMLResponse)
def login_get(request: Request):
    return templates.TemplateResponse(request=request, name="login.html", context={"request": request})

@app.post("/login")
def login_post(
    request: Request,
    username: str = Form(...),
    password: str = Form(...),
    db: Session = Depends(get_db)
):
    user = db.query(models.User).filter(models.User.username == username).first()
    if not user or not verify_password(password, user.password_hash):
        return templates.TemplateResponse(request=request, name="login.html", context={"request": request, "error": "Credenciales inválidas"})
    
    access_token = create_access_token(data={"sub": user.username})
    
    response = RedirectResponse(url="/", status_code=status.HTTP_302_FOUND)
    response.set_cookie(key="access_token", value=f"{access_token}", httponly=True)
    return response

@app.get("/logout")
def logout():
    response = RedirectResponse(url="/login", status_code=status.HTTP_302_FOUND)
    response.delete_cookie("access_token")
    return response

EMPLEADOS_DIRECTORIO = {
    "AGUIRRE JAIMES MARIA ALEJANDRA": {"cargo": "COORDINADORA INVESTIGACIÓN Y DESARROLLO DE PRODUCTO", "celular": "3226372036"},
    "ARISTIZABAL ECHEVERRY LUIS FELIPE": {"cargo": "GERENTE", "celular": "3174400191"},
    "ARROYAVE PEREZ DEILY DAYHANA": {"cargo": "ASESOR COMERCIAL", "celular": "3226040898"},
    "BERNAL VARGAS MARIA JOSÉ": {"cargo": "APRENDIZ SENA ETAPA PRODUCTIVA", "celular": "3116939284"},
    "CAMPILLO VELASQUEZ LUISA MARIA": {"cargo": "ANALISTA JUNIOR DISEÑO CREATIVO", "celular": "3216507532"},
    "CARDONA OSORIO LAURA": {"cargo": "LIDER C&DH", "celular": "3005068231"},
    "CORREA RESTREPO LEIDY JHOANA": {"cargo": "APRENDIZ SENA ETAPA PRODUCTIVA", "celular": "3195229166"},
    "CORRALES MARÍN JUAN ESTEBAN": {"cargo": "COORDINADOR FINANCIERO", "celular": "3104591966"},
    "FORERO VELÁZQUEZ JESSICA": {"cargo": "ASESOR COMERCIAL", "celular": "3113821193"},
    "GAVIRIA CASTAÑO SOPHIA": {"cargo": "ANALISTA JUNIOR DISEÑO CREATIVO", "celular": "3235257501"},
    "GARCIA CARDONA LUISA FERNANDA": {"cargo": "ANALISTA JUNIOR CALIDAD Y CONFECCION", "celular": "3127743024"},
    "GIRALDO HERRERA CRISTIAN FELIPE": {"cargo": "LIDER DESARROLLO ORGANIZACIONAL", "celular": "3215652399"},
    "HERNANDEZ FRANCO DAIRO": {"cargo": "AUXILIAR DISEÑO TEJEDURIA", "celular": "3207279081"},
    "MERCADO RENDON YEIMMY ALEJANDRA": {"cargo": "SUPERVISORA DE PLANTA", "celular": "3202164992"},
    "OCAMPO CANDAMIL DIEGO ALEJANDRO": {"cargo": "AUXILIAR COMERCIAL", "celular": "3225401861"},
    "OSORIO MANSO YICETH CAMILA": {"cargo": "APRENDIZ SENA ETAPA PRODUCTIVA", "celular": "3116857931"},
    "PALACIO HERRERA CARLOS": {"cargo": "COORDINADOR DE PRODUCCIÓN", "celular": "3206041907"},
    "PARRA GALVIS ADRIANA": {"cargo": "COORDINADOR DE PRODUCCIÓN", "celular": "3136334441"},
    "PEREZ CIFUENTES MARIA ALEJANDRA": {"cargo": "AUXILIAR DE DISEÑO", "celular": "3043749051"},
    "QUINTERO TAPASCO ANGIE DANIELA": {"cargo": "ANALISTA JUNIOR CALIDAD Y CONFECCION", "celular": "3234623428"},
    "QUINTERO LONDOÑO CRISTIAN DAVID": {"cargo": "SUPERVISOR DE PLANTA", "celular": "3203367730"},
    "SANCHEZ ORTEGA KAREN NAYITH": {"cargo": "ANALISTA DE SEGURIDAD Y SALUD EN EL TRABAJO", "celular": "3226124959"},
    "SALAZAR ARISTIZABAL ALEX MAURICIO": {"cargo": "AUXILIAR CONTABLE", "celular": "3045916546"},
    "SALAMANCA VILLA ANGIE JULIETH": {"cargo": "COORDINADORA JUNIOR COMERCIAL", "celular": "3116125157"},
    "TEJADA PEREZ PAULINA": {"cargo": "ANALISTA JUNIOR CALIDAD Y CONFECCION", "celular": "3234908094"},
    "TORO GARCIA JUAN JOSE": {"cargo": "COORDINADOR LOGISTICO", "celular": "3186972893"},
    "TORRES MONTAÑO NICOLL XIOMARA": {"cargo": "AUXILIAR DE TESORERIA", "celular": "3112006100"},
    "VARGAS CORREA DIANA MARCELA": {"cargo": "ANALISTA DE METODOS Y EFICIENCIA PRODUCTIVA", "celular": "3128944490"},
    "VELEZ FLOREZ VALENTINA": {"cargo": "COORDINADORA NUEVAS CUENTAS", "celular": "3168316404"},
    "VELEZ JIMENEZ JUAN CAMILO": {"cargo": "ANALISTA JUNIOR METODOS, TIEMPO, CALIDAD", "celular": "3053027288"},
    "ZAPATA AGUIRRE MARIA ALEJANDRA": {"cargo": "ANALISTA ADMINISTRATIVA", "celular": "3245646745"}
}

@app.get("/tickets", response_class=HTMLResponse)
def client_dashboard(request: Request, user: models.User = Depends(get_current_user), db: Session = Depends(get_db)):
    if not user:
        return RedirectResponse(url="/login", status_code=status.HTTP_302_FOUND)
    
    # Mostrar últimos tickets
    tickets = db.query(models.Ticket).order_by(models.Ticket.fecha_creacion.desc()).all()
    for t in tickets:
        if not t.celular_solicitante and t.nombre_solicitante in EMPLEADOS_DIRECTORIO:
            t.celular_solicitante = EMPLEADOS_DIRECTORIO[t.nombre_solicitante]["celular"]

    return templates.TemplateResponse(request=request, name="cliente_dashboard.html", context={"request": request, "tickets": tickets, "user": user})

@app.post("/tickets")
def create_ticket(
    request: Request,
    titulo: str = Form(...),
    descripcion: str = Form(...),
    nombre_solicitante: str = Form(...),
    cargo_solicitante: str = Form(None),
    celular_solicitante: str = Form(None),
    urgencia: str = Form(...),
    tipo_afectacion: str = Form(...),
    subtipo_equipo: str = Form(None),
    tipo_solicitud: str = Form(...),
    archivos: List[UploadFile] = File(None),
    user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    if not user:
        return RedirectResponse(url="/login", status_code=status.HTTP_302_FOUND)
    
    # Si no llegó el celular, autocompletar desde el directorio
    if not celular_solicitante and nombre_solicitante in EMPLEADOS_DIRECTORIO:
        celular_solicitante = EMPLEADOS_DIRECTORIO[nombre_solicitante]["celular"]

    new_ticket = models.Ticket(
        titulo=titulo,
        descripcion=descripcion,
        nombre_solicitante=nombre_solicitante,
        cargo_solicitante=cargo_solicitante,
        celular_solicitante=celular_solicitante,
        urgencia=urgencia,
        tipo_afectacion=tipo_afectacion,
        subtipo_equipo=subtipo_equipo,
        tipo_solicitud=tipo_solicitud,
        responsable="Soporte" # Por defecto entra a Soporte
    )
    db.add(new_ticket)
    db.commit()
    db.refresh(new_ticket)

    # Guardar archivos adjuntos si existen (máximo 3)
    if archivos:
        for f in archivos[:3]:
            if f.filename and f.filename.strip():
                contenido = f.file.read()
                if len(contenido) > 0:
                    adjunto = models.TicketAdjunto(
                        ticket_id=new_ticket.id,
                        nombre_archivo=f.filename,
                        tipo_contenido=f.content_type or "application/octet-stream",
                        archivo_bytes=contenido,
                        tamano=len(contenido)
                    )
                    db.add(adjunto)
        db.commit()

    return RedirectResponse(url="/tickets", status_code=status.HTTP_302_FOUND)

@app.get("/admin", response_class=HTMLResponse)
def admin_dashboard(request: Request, user: models.User = Depends(get_current_admin), db: Session = Depends(get_db)):
    if not user:
        return RedirectResponse(url="/login", status_code=status.HTTP_302_FOUND)
    
    tickets = db.query(models.Ticket).order_by(models.Ticket.fecha_creacion.desc()).all()
    for t in tickets:
        if not t.celular_solicitante and t.nombre_solicitante in EMPLEADOS_DIRECTORIO:
            t.celular_solicitante = EMPLEADOS_DIRECTORIO[t.nombre_solicitante]["celular"]

    return templates.TemplateResponse(request=request, name="admin_dashboard.html", context={"request": request, "tickets": tickets, "user": user})

@app.post("/admin/ticket/{ticket_id}/status")
def update_ticket_status(
    request: Request,
    ticket_id: int,
    estado: str = Form(...),
    responsable: str = Form(...),
    user: models.User = Depends(get_current_admin),
    db: Session = Depends(get_db)
):
    if not user:
        return RedirectResponse(url="/login", status_code=status.HTTP_302_FOUND)
    
    ticket = db.query(models.Ticket).filter(models.Ticket.id == ticket_id).first()
    if ticket:
        ticket.estado = estado
        ticket.responsable = responsable
        db.commit()
        
    referer = request.headers.get("referer", "/admin")
    return RedirectResponse(url=referer, status_code=status.HTTP_302_FOUND)

@app.get("/admin/ticket/{ticket_id}", response_class=HTMLResponse)
def admin_ticket_detail(request: Request, ticket_id: int, user: models.User = Depends(get_current_admin), db: Session = Depends(get_db)):
    if not user:
        return RedirectResponse(url="/login", status_code=status.HTTP_302_FOUND)
    
    ticket = db.query(models.Ticket).filter(models.Ticket.id == ticket_id).first()
    if not ticket:
        raise HTTPException(status_code=404, detail="Ticket no encontrado")
    
    if not ticket.celular_solicitante and ticket.nombre_solicitante in EMPLEADOS_DIRECTORIO:
        ticket.celular_solicitante = EMPLEADOS_DIRECTORIO[ticket.nombre_solicitante]["celular"]
        
    return templates.TemplateResponse(request=request, name="admin_ticket_detail.html", context={"request": request, "ticket": ticket, "user": user})

@app.post("/admin/ticket/{ticket_id}/gestion")
def add_ticket_gestion(
    ticket_id: int,
    nota: str = Form(...),
    user: models.User = Depends(get_current_admin),
    db: Session = Depends(get_db)
):
    if not user:
        return RedirectResponse(url="/login", status_code=status.HTTP_302_FOUND)
    
    ticket = db.query(models.Ticket).filter(models.Ticket.id == ticket_id).first()
    if ticket:
        gestion = models.TicketGestion(
            ticket_id=ticket.id,
            autor=user.username,
            nota=nota
        )
        db.add(gestion)
        # Cambiamos estado automáticamente a En Progreso si estaba Pendiente
        if ticket.estado == "Pendiente":
            ticket.estado = "En Progreso"
        db.commit()
        
    return RedirectResponse(url=f"/admin/ticket/{ticket_id}", status_code=status.HTTP_302_FOUND)

@app.post("/admin/ticket/{ticket_id}/delete")
def delete_ticket(
    ticket_id: int,
    user: models.User = Depends(get_current_admin),
    db: Session = Depends(get_db)
):
    if not user:
        return RedirectResponse(url="/login", status_code=status.HTTP_302_FOUND)
    
    ticket = db.query(models.Ticket).filter(models.Ticket.id == ticket_id).first()
    if ticket:
        db.delete(ticket)
        db.commit()
        
    return RedirectResponse(url="/admin", status_code=status.HTTP_302_FOUND)

@app.get("/attachments/{adjunto_id}")
def get_attachment(adjunto_id: int, db: Session = Depends(get_db)):
    adjunto = db.query(models.TicketAdjunto).filter(models.TicketAdjunto.id == adjunto_id).first()
    if not adjunto:
        raise HTTPException(status_code=404, detail="Archivo no encontrado")
    return Response(
        content=adjunto.archivo_bytes,
        media_type=adjunto.tipo_contenido or "application/octet-stream",
        headers={"Content-Disposition": f'inline; filename="{adjunto.nombre_archivo}"'}
    )

@app.get("/admin/export-excel")
def export_tickets_excel(user: models.User = Depends(get_current_admin), db: Session = Depends(get_db)):
    if not user:
        return RedirectResponse(url="/login", status_code=status.HTTP_302_FOUND)
    
    tickets = db.query(models.Ticket).order_by(models.Ticket.fecha_creacion.desc()).all()
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reporte Tickets Olamtex"
    
    # Estilos profesionales
    header_fill = PatternFill(start_color="312E81", end_color="312E81", fill_type="solid") # Indigo 900
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    row_font = Font(name="Segoe UI", size=10)
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    
    headers = [
        "ID Ticket", "Fecha Creación", "Solicitante", "Cargo", "Celular / WhatsApp",
        "Título del Problema", "Afectación", "Subtipo Equipo", "Tipo Solicitud",
        "Responsable Asignado", "Nivel Urgencia", "Estado Actual", "Descripción del Problema", 
        "Historial de Gestiones y Notas", "Archivos Adjuntos"
    ]
    
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border
    ws.row_dimensions[1].height = 28
    
    for t in tickets:
        cel = t.celular_solicitante or (EMPLEADOS_DIRECTORIO.get(t.nombre_solicitante, {}).get("celular", ""))
        gestiones_str = " | ".join([f"[{g.fecha.strftime('%d/%m/%Y %H:%M')} - {g.autor}]: {g.nota}" for g in t.gestiones]) if t.gestiones else "Sin gestiones registradas"
        num_adjuntos = f"{len(t.adjuntos)} archivo(s)" if t.adjuntos else "Sin archivos"
        
        row_data = [
            t.id,
            t.fecha_creacion.strftime('%d/%m/%Y %I:%M %p') if t.fecha_creacion else "",
            t.nombre_solicitante,
            t.cargo_solicitante or "",
            cel,
            t.titulo,
            t.tipo_afectacion,
            t.subtipo_equipo or "N/A",
            t.tipo_solicitud,
            t.responsable,
            t.urgencia,
            t.estado,
            t.descripcion,
            gestiones_str,
            num_adjuntos
        ]
        ws.append(row_data)
        current_row = ws.max_row
        for cell in ws[current_row]:
            cell.font = row_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[current_row].height = 22

    # Autoajustar ancho de columnas
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(min(max_len + 4, 60), 14)

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    
    filename = f"Reporte_Tickets_Olamtex_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return Response(
        content=stream.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )

